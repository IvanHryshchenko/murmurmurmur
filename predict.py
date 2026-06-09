import os
import datetime
import numpy as np
from openpyxl import load_workbook

from data_preparation import (
    load_cutting_sheet,
    load_all_filters,
    compute_detailed_score,
    create_cutting_features,
    _encode_machine_group,
    FEATURE_COLS_CUTTING,
    GENERIC_SHEETS,
    SIMPLE_QTY_SHEETS,
    load_sheet_dynamic,
    load_and_prepare_total_gcsd,
    create_features_total_gcsd,
    FEATURE_COLS_TOTAL_GCSD,
    _is_header, _to_float, _col_index,
)
from model import MiniAI

FILE_PATH  = 'K0_old.xlsx'
MODELS_DIR = 'models'
OUTPUT_DIR = 'results'
NUMBER_FMT = '0.0000'

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs('logs', exist_ok=True)

_log_file = open(os.path.join('logs', 'predict.log'), 'a', encoding='utf-8')


def log(text=''):
    print(text)
    _log_file.write(text + '\n')
    _log_file.flush()


def safe_name(s):
    return s.replace(' ', '_').replace('/', '_').replace('\n', '_')


def model_path(sheet_name):
    return os.path.join(MODELS_DIR, safe_name(sheet_name), 'model.pkl')


def _round4(v):
    return round(float(v), 4)


def _write_cell(cell, value, fmt=NUMBER_FMT):
    cell.value = _round4(value)
    cell.number_format = fmt


# ═══════════════════════════════════════════════════════════════════════════════

session_start = datetime.datetime.now()
W = 70

log()
log('  ' + '═' * W)
log(f'  {"🔮  BEST WIRE PREDICTION (Neural Network — Detailed Score)":^{W}}')
log(f'  {"Started: " + session_start.strftime("%Y-%m-%d  %H:%M:%S"):^{W}}')
log('  ' + '═' * W)

all_results = []
wb = load_workbook(FILE_PATH)


# ─── 1. CUTTING — поиск лучшего провода ──────────────────────────────────────

log()
log('  ' + '=' * W)
log('  CUTTING — нейросеть предсказывает детальный score для каждого провода')
log('  ' + '=' * W)
log()

df_cutting = load_cutting_sheet(FILE_PATH, 'CUTTING')
log(f'  Загружено {len(df_cutting)} проводов из CUTTING')

# Загружаем фильтры (для верификации реального score)
log()
log('  Загрузка фильтров операций...')
filters_dict = load_all_filters(FILE_PATH)
total_ops = sum(len(v) for v in filters_dict.values())
log(f'  Итого операций с ограничениями: {total_ops}')

# Вычисляем реальный score для каждого провода (верификация)
real_scores = []
real_lp = []
real_lpfa = []
real_hv = []
for _, w in df_cutting.iterrows():
    mid_g = (w['min_gage'] + w['max_gage']) / 2.0
    mid_l = (w['min_len']  + w['max_len'])  / 2.0
    total, s_lp, s_lpfa, s_hv = compute_detailed_score(mid_g, mid_l, filters_dict)
    real_scores.append(total)
    real_lp.append(s_lp)
    real_lpfa.append(s_lpfa)
    real_hv.append(s_hv)

real_scores = np.array(real_scores, dtype=float)

log()
log(f'  Реальный score: min={real_scores.min():.0f}  max={real_scores.max():.0f}'
    f'  mean={real_scores.mean():.1f}')

# ─── Нейросеть ───────────────────────────────────────────────────────────────

mp     = model_path('CUTTING')
use_ai = os.path.exists(mp)

if use_ai:
    ai_cut = MiniAI.load(mp)

    if hasattr(ai_cut, 'group_mapping') and ai_cut.group_mapping:
        gmap, ggmean = ai_cut.group_mapping, ai_cut.group_global_mean or 0.0
    else:
        _, gmap, ggmean = _encode_machine_group(df_cutting)

    df_feat = create_cutting_features(df_cutting, gmap, ggmean)
    X       = df_feat[FEATURE_COLS_CUTTING].values

    nn_scores = ai_cut.predict(X)

    mae = float(np.mean(np.abs(nn_scores - real_scores)))
    log()
    log(f'  ✅ Нейросеть предсказала score для {len(nn_scores)} проводов')
    log(f'     mean={nn_scores.mean():.2f}  min={nn_scores.min():.2f}  max={nn_scores.max():.2f}')
    log(f'     MAE vs реальный score: {mae:.3f}')
else:
    # Fallback: используем реальный score напрямую
    log('  ⚠️  Модель не найдена — fallback (реальный score из правил)')
    nn_scores = real_scores.copy()

# ─── ВЫБОР ЛУЧШЕГО ПРОВОДА ────────────────────────────────────────────────────
# Критерий: max(predicted_score), при равенстве — min(gcsp)

gcsp_arr = df_cutting['gcsp'].values
ranking  = nn_scores * 10000.0 - gcsp_arr   # × 10000 гарантирует приоритет score над gcsp
best_idx = int(np.argmax(ranking))
best_row = df_cutting.iloc[best_idx]

# Детализация реального score для лучшего провода
best_mid_g = (best_row['min_gage'] + best_row['max_gage']) / 2
best_mid_l = (best_row['min_len']  + best_row['max_len'])  / 2
_, best_lp, best_lpfa, best_hv = compute_detailed_score(
    best_mid_g, best_mid_l, filters_dict)

log()
log('  ' + '─' * W)
log(f'  {"🏆  ЛУЧШИЙ ПРОВОД":^{W}}')
log('  ' + '─' * W)
log(f'  Excel-строка      : {int(best_row["excel_row"])}')
log(f'  Группа машины     : {best_row["machine_group"]}')
log(f'  Gage              : [{best_row["min_gage"]:.2f} – {best_row["max_gage"]:.2f}]  (mid={best_mid_g:.2f})')
log(f'  Length            : [{best_row["min_len"]:.0f} – {best_row["max_len"]:.0f}]  (mid={best_mid_l:.0f})')
log(f'  GCSP              : {best_row["gcsp"]:.4f}  сек/шт')
log()
log(f'  NN predicted score: {nn_scores[best_idx]:.2f}')
log(f'  Реальный score    : {int(real_scores[best_idx])}  из {total_ops} операций')
log()
log(f'  Детализация реального score:')
log(f'    LEAD PREP:    {best_lp:4d}  из {len(filters_dict["LEAD PREP"])}  операций')
log(f'    LEAD PREP FA: {best_lpfa:4d}  из {len(filters_dict["LEAD PREP FA"])} операций')
log(f'    High Voltage: {best_hv:4d}  из {len(filters_dict["High Voltage"])}  операций')

# ─── Топ-10 ───────────────────────────────────────────────────────────────────

log()
log('  Топ-10 проводов по предсказанию нейросети:')
log(f'  {"Строка":>6} {"Gage мин":>8} {"Gage макс":>9} {"Len мин":>7} {"Len макс":>8}'
    f' {"GCSP":>7} {"NN score":>9} {"Real":>6}')
log('  ' + '─' * 62)
top10_idx = np.argsort(ranking)[::-1][:10]
for idx in top10_idx:
    r = df_cutting.iloc[idx]
    log(f'  {int(r["excel_row"]):>6} {r["min_gage"]:>8.2f} {r["max_gage"]:>9.2f}'
        f' {r["min_len"]:>7.0f} {r["max_len"]:>8.0f}'
        f' {r["gcsp"]:>7.4f} {nn_scores[idx]:>9.2f} {int(real_scores[idx]):>6}')

# ─── Записываем в Excel ───────────────────────────────────────────────────────
# Col I = TOTAL Qty Leads — единственный входной столбец;
# Col J (SUB TOTAL = I*H) и Col L (TOTAL TIME = J) — формулы Excel, не трогаем.
# TOTAL values GCSD ссылается на итоговые ячейки J685/L685/L687 через формулы,
# поэтому достаточно правильно заполнить Col I.

ws_cut = wb['CUTTING']
written_cut = 0

# Обнуляем QTY (col I) для всех проводов из df_cutting
for idx, (_, row_data) in enumerate(df_cutting.iterrows()):
    excel_row = int(row_data['excel_row'])
    cell_i = ws_cut.cell(row=excel_row, column=9)
    # Сбрасываем только если там не формула
    if not (isinstance(cell_i.value, str) and cell_i.value.startswith('=')):
        cell_i.value = 0
        cell_i.number_format = '0'
    written_cut += 1

# Лучший провод: QTY=1 в столбец I, отметка BEST в столбец K (Qty Marked Leads)
best_excel_row = int(best_row['excel_row'])
ws_cut.cell(row=best_excel_row, column=9).value  = 1
ws_cut.cell(row=best_excel_row, column=9).number_format = '0'
ws_cut.cell(row=best_excel_row, column=11).value = 'BEST'

log()
log(f'  ✅ QTY проставлен для {written_cut} строк (0 для всех, 1 для лучшего).')
log(f'     Лучший провод: QTY=1 (col I) + BEST (col K) → строка {best_excel_row}')
log(f'     TOTAL values GCSD обновится через формулы CUTTING!J685/L685/L687.')

all_results.append(dict(
    sheet='CUTTING', rows=len(df_cutting), written=written_cut,
    method='AI' if use_ai else 'fallback',
    best_wire_row=best_excel_row,
    best_nn_score=float(nn_scores[best_idx]),
    best_real_score=int(real_scores[best_idx]),
))


# ─── 2. TOTAL VALUES GCSD ──────────────────────────────────────────────────────
# Этот лист содержит формулы, ссылающиеся напрямую на итоговые ячейки CUTTING
# (J685, L685, L687) и других листов. После заполнения col I в CUTTING
# значения пересчитаются автоматически при открытии в Excel.
# Восстанавливаем только формулы SUM на случай если они были перезаписаны ранее.

log()
log('  ' + '=' * W)
log('  Sheet: TOTAL values GCSD')
log('  ' + '=' * W)

ws_g   = wb['TOTAL values GCSD']
E_COL, M_COL = 5, 13

# Восстанавливаем формулы итогов (на случай повреждения предыдущими запусками)
ws_g.cell(row=13, column=E_COL).value = '=SUM(E6:E12)'
ws_g.cell(row=23, column=E_COL).value = '=SUM(E19:E22)'
ws_g.cell(row=27, column=M_COL).value = '=SUM(M19:M26)'

log(f'  ✅ Формулы SUM восстановлены. Значения CUTTING → TOTAL values GCSD')
log(f'     обновятся автоматически через формулы при открытии Excel.')
all_results.append(dict(sheet='TOTAL values GCSD', rows=0,
                        written=3, method='formulas',
                        best_wire_row=None, best_nn_score=None, best_real_score=None))


# ─── 3 + 4. QTY=1 для всех рабочих листов ──────────────────────────────────
#
# Правило одинаково для всех 6 листов:
#   QTY = 1  если:  числовой GCSP > 0  И  нет пометки "не учитываю"
#   QTY = 0  иначе  (явный 0 чтобы формула K=J*H не давала пустоту)
#
# "не учитываю" встречается только в LEAD PREP (col N = 14).
# LEAD PREP FA идентичен по структуре, но пометки там нет.
# Для остальных листов (High Voltage, FA Conns and wires, FA Taping,
# FA Miscellaneos) пометки тоже нет — просто проверяем числовой GCSP.
#
# Колонки K/TOTAL содержат формулы =J*H — НЕ трогаем их.
# TOTAL values GCSD получает значения через цепочку формул автоматически.

from openpyxl.cell.cell import MergedCell

def _is_merged(cell):
    return isinstance(cell, MergedCell)

# Конфигурация каждого листа:
#   gcsp_col  — 1-based колонка GLOBAL SEC/PC (числовое время)
#   qty_col   — 1-based колонка QTY (куда пишем 0 или 1)
#   header_row — строка-заголовок (данные начинаются со следующей)
#   note_col  — 1-based колонка с "не учитываю" (None если нет)

QTY_SHEETS_CFG = [
    ('LEAD PREP',          {'gcsp_col':  8, 'qty_col': 10, 'header_row': 2, 'note_col': 14}),
    ('LEAD PREP FA',       {'gcsp_col':  8, 'qty_col': 10, 'header_row': 2, 'note_col': 14}),
    ('High Voltage',       {'gcsp_col':  8, 'qty_col': 10, 'header_row': 4, 'note_col': None}),
    ('FA Conns and wires', {'gcsp_col': 17, 'qty_col': 19, 'header_row': 3, 'note_col': None}),
    ('FA Taping',          {'gcsp_col':  7, 'qty_col':  9, 'header_row': 2, 'note_col': None}),
    ('FA Miscellaneos',    {'gcsp_col':  8, 'qty_col': 10, 'header_row': 2, 'note_col': None}),
]

for sheet_name, cfg in QTY_SHEETS_CFG:
    log()
    log(f'  Sheet: {sheet_name}  [QTY=1 mode]')

    if sheet_name not in wb.sheetnames:
        log(f'  ⚠️  Лист не найден, пропуск.')
        all_results.append(dict(sheet=sheet_name, rows=0, written=0,
                                method='QTY=1', best_wire_row=None,
                                best_nn_score=None, best_real_score=None))
        continue

    ws         = wb[sheet_name]
    gcsp_col   = cfg['gcsp_col']
    qty_col    = cfg['qty_col']
    header_row = cfg['header_row']
    note_col   = cfg['note_col']

    qty_written = qty_zeroed = 0

    for row_cells in ws.iter_rows(min_row=header_row + 1):
        rnum = row_cells[0].row

        # ── Пропускаем полностью пустые строки (за границей данных) ──────────
        row_has_data = any(
            not _is_merged(c) and c.value is not None
            for c in row_cells
            if not (qty_col <= len(row_cells) and c.column == qty_col)
        )
        if not row_has_data:
            continue

        # ── Ячейка QTY ────────────────────────────────────────────────────────
        if qty_col > len(row_cells):
            continue
        qcell = row_cells[qty_col - 1]
        if _is_merged(qcell):
            continue
        if isinstance(qcell.value, str) and qcell.value.startswith('='):
            continue   # формула-агрегат — не трогаем

        # ── Проверяем GCSP ────────────────────────────────────────────────────
        if gcsp_col > len(row_cells):
            qcell.value = 0
            continue
        gcsp_cell = row_cells[gcsp_col - 1]
        gcsp_val  = None if _is_merged(gcsp_cell) else gcsp_cell.value
        has_gcsp  = isinstance(gcsp_val, (int, float)) and gcsp_val > 0

        # ── Проверяем пометку "не учитываю" ──────────────────────────────────
        skip = False
        if note_col and note_col <= len(row_cells):
            nc = row_cells[note_col - 1]
            if not _is_merged(nc) and nc.value:
                skip = 'не учитываю' in str(nc.value).lower()

        # ── Записываем ───────────────────────────────────────────────────────
        if has_gcsp and not skip:
            qcell.value = 1
            qcell.number_format = '0'
            qty_written += 1
        else:
            qcell.value = 0
            qcell.number_format = '0'
            qty_zeroed += 1

    log(f'  ✅ QTY=1: {qty_written} строк,  QTY=0: {qty_zeroed} строк')
    all_results.append(dict(sheet=sheet_name, rows=qty_written + qty_zeroed,
                            written=qty_written, method='QTY=1',
                            best_wire_row=None, best_nn_score=None, best_real_score=None))


# ─── СОХРАНЕНИЕ И ИТОГ ────────────────────────────────────────────────────────

combined_out = os.path.join(OUTPUT_DIR, 'RESULT_ALL_SHEETS.xlsx')
wb.save(combined_out)

session_elapsed = (datetime.datetime.now() - session_start).total_seconds()

log()
log('  ' + '═' * W)
log(f'  {"📋  SESSION SUMMARY":^{W}}')
log('  ' + '═' * W)
log()

for r in all_results:
    bw = (f'row {r["best_wire_row"]}  NN={r["best_nn_score"]:.2f}  real={r["best_real_score"]}'
          if r['best_wire_row'] is not None else '—')
    log(f'  {r["sheet"]:<28}  method={r["method"]:10}  best={bw}')

# ─── ЛУЧШИЙ ПРОВОД — финальный баннер ────────────────────────────────────────
log()
log('  ' + '═' * W)
log(f'  {"🏆  ЛУЧШИЙ ПРОВОД  🏆":^{W}}')
log('  ' + '═' * W)
log(f'  {"Excel-строка":<22}: {best_excel_row}')
log(f'  {"Группа машины":<22}: {best_row["machine_group"]}')
log(f'  {"Gage":<22}: [{best_row["min_gage"]:.2f} – {best_row["max_gage"]:.2f}]  (mid={best_mid_g:.2f})')
log(f'  {"Length":<22}: [{best_row["min_len"]:.0f} – {best_row["max_len"]:.0f}]  (mid={best_mid_l:.0f})')
log(f'  {"GCSP":<22}: {best_row["gcsp"]:.4f}  сек/шт')
log()
log(f'  {"NN predicted score":<22}: {nn_scores[best_idx]:.2f}')
log(f'  {"Реальный score":<22}: {int(real_scores[best_idx])}  из {total_ops} операций')
log()
log(f'  Детализация по пространствам:')
lp_total   = len(filters_dict["LEAD PREP"])
lpfa_total = len(filters_dict["LEAD PREP FA"])
hv_total   = len(filters_dict["High Voltage"])
log(f'    {"LEAD PREP":<20}: {best_lp:3d}  из {lp_total:2d}  операций  ({100*best_lp//max(1,lp_total)}%)')
log(f'    {"LEAD PREP FA":<20}: {best_lpfa:3d}  из {lpfa_total:2d}  операций  ({100*best_lpfa//max(1,lpfa_total)}%)')
log(f'    {"High Voltage":<20}: {best_hv:3d}  из {hv_total:2d}  операций  ({100*best_hv//max(1,hv_total)}%)')
log()
log(f'  📁  Результат сохранён: {combined_out}')
log(f'  ⏱  Общее время:         {session_elapsed:.1f} сек')
log('  ' + '═' * W)

_log_file.close()