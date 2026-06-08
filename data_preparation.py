import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook


# ─── ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ─────────────────────────────────────────────────

def _to_float(v):
    try:
        f = float(v)
        return None if (np.isnan(f) or np.isinf(f)) else f
    except (TypeError, ValueError):
        return None


def _is_ne(row_vals):
    """Пропускаем строки 'не учитываю'."""
    for v in row_vals:
        if v is not None and 'не учитываю' in str(v).lower():
            return True
    return False


def _is_header(row_vals):
    """Строка является заголовком, если содержит 'global sec'."""
    for v in row_vals:
        if v is not None and 'global sec' in str(v).lower():
            return True
    return False


def _col_index(row_vals, *keywords):
    for kw in keywords:
        for j, v in enumerate(row_vals):
            if v is not None and kw.lower() in str(v).lower():
                return j
    return None


def _parse_mm2(val):
    """'0.5 mm²' / 30mm² / 0.5 → float."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return _to_float(val)
    s = str(val).strip().lower().replace('mm²', '').replace('mm2', '').strip()
    try:
        return float(s)
    except ValueError:
        return None


# ─── CUTTING ──────────────────────────────────────────────────────────────────

def load_cutting_sheet(file_path, sheet_name='CUTTING'):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows_data = list(ws.iter_rows(min_row=1, values_only=True))

    COL_MIN_GAGE = 3
    COL_MAX_GAGE = 4
    COL_MIN_LEN  = 5
    COL_MAX_LEN  = 6
    COL_GCSP     = 7
    COL_QTY      = 8

    records = []
    current_group = 'Unknown'

    for i in range(2, min(681, len(rows_data))):
        row = rows_data[i]

        comment_val = row[0] if len(row) > 0 else None
        if comment_val is not None:
            s = str(comment_val).strip()
            if s and len(s) > 2:
                current_group = s[:80]

        def _get(col):
            return row[col] if len(row) > col else None

        gcsp = _to_float(_get(COL_GCSP))
        if gcsp is None or gcsp <= 0:
            continue

        min_gage = _to_float(_get(COL_MIN_GAGE))
        max_gage = _to_float(_get(COL_MAX_GAGE))
        min_len  = _to_float(_get(COL_MIN_LEN))
        max_len  = _to_float(_get(COL_MAX_LEN))

        has_gage = (min_gage is not None or max_gage is not None)
        has_len  = (min_len  is not None or max_len  is not None)
        if not (has_gage or has_len):
            continue

        records.append({
            'row_index':     i,
            'excel_row':     i + 1,
            'machine_group': current_group,
            'min_gage':      min_gage if min_gage is not None else 0.0,
            'max_gage':      max_gage if max_gage is not None else 0.0,
            'min_len':       min_len  if min_len  is not None else 0.0,
            'max_len':       max_len  if max_len  is not None else 0.0,
            'gcsp':          gcsp,
        })

    wb.close()
    df = pd.DataFrame(records) if records else pd.DataFrame(columns=[
        'row_index', 'excel_row', 'machine_group',
        'min_gage', 'max_gage', 'min_len', 'max_len', 'gcsp'
    ])
    return df


def _encode_machine_group(df_cutting):
    log_gcsp = np.log1p(df_cutting['gcsp'])
    mean_per_group = log_gcsp.groupby(df_cutting['machine_group']).mean()
    global_mean = float(log_gcsp.mean())
    mapping = mean_per_group.to_dict()
    encoded = df_cutting['machine_group'].map(mapping).fillna(global_mean)
    return encoded, mapping, global_mean


# ─── ЗАГРУЗКА ФИЛЬТРОВ ────────────────────────────────────────────────────────
#
# Каждая строка фильтра = одна операция с ограничениями по gage и/или length.
# Провод «проходит» операцию, если его параметры попадают в ограничения.
# Строки без ограничений (min_gage=None, max_gage=None, min_len=None, max_len=None)
# НЕ учитываются в score, так как они всегда проходят и не дают информации.
#
# score = sum(проходит ли провод каждую операцию) по всем трём пространствам

def _load_filter_rows(file_path, sheet_name, gage_cols=(4, 5), len_cols=None):
    """
    Загружает операции из листа как набор ограничений.
    Возвращает список dict: {min_gage, max_gage, min_len, max_len, gcsp}
    Включает ТОЛЬКО строки у которых есть хотя бы одно ограничение.
    """
    wb  = load_workbook(file_path, read_only=True, data_only=True)
    ws  = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    gc, gc2 = gage_cols
    lc1 = len_cols[0] if len_cols else None
    lc2 = len_cols[1] if len_cols else None

    result = []
    for row in rows:
        vals = list(row)
        if _is_ne(vals):
            continue

        gcsp = _to_float(vals[7]) if len(vals) > 7 else None
        if gcsp is None or gcsp <= 0:
            continue

        min_g = _parse_mm2(vals[gc]  if len(vals) > gc  else None)
        max_g = _parse_mm2(vals[gc2] if len(vals) > gc2 else None)
        min_l = _to_float(vals[lc1]  if lc1 is not None and len(vals) > lc1 else None)
        max_l = _to_float(vals[lc2]  if lc2 is not None and len(vals) > lc2 else None)

        # Берём только строки с реальными ограничениями
        has_gage = (min_g is not None or max_g is not None)
        has_len  = (min_l is not None or max_l is not None)
        if not (has_gage or has_len):
            continue

        result.append({
            'min_gage': min_g, 'max_gage': max_g,
            'min_len':  min_l, 'max_len':  max_l,
            'gcsp':     gcsp,
        })

    return result


def load_all_filters(file_path):
    """
    Загружает ограничения из трёх рабочих пространств.
    LEAD PREP / LEAD PREP FA: ограничения по gage (col 4, 5)
    High Voltage: ограничения по length (col 5, 6) И по gage в формате mm² (col 5, 6)
    """
    filters = {
        'LEAD PREP':    _load_filter_rows(file_path, 'LEAD PREP',    gage_cols=(4, 5)),
        'LEAD PREP FA': _load_filter_rows(file_path, 'LEAD PREP FA', gage_cols=(4, 5)),
        'High Voltage': _load_filter_rows(file_path, 'High Voltage', gage_cols=(5, 6), len_cols=(5, 6)),
    }
    for name, flist in filters.items():
        print(f'  Фильтры {name}: {len(flist)} операций с ограничениями')
    return filters


def wire_passes_operation(mid_gage, mid_len, op):
    """
    Проверяет: провод с (mid_gage, mid_len) проходит ли ограничения операции.
    Ограничения по gage и length независимы — оба должны выполняться.
    """
    # Ограничения по gage
    if op['min_gage'] is not None and mid_gage < op['min_gage']:
        return False
    if op['max_gage'] is not None and mid_gage > op['max_gage']:
        return False
    # Ограничения по length
    if op['min_len'] is not None and mid_len < op['min_len']:
        return False
    if op['max_len'] is not None and mid_len > op['max_len']:
        return False
    return True


def compute_detailed_score(mid_gage, mid_len, filters_dict):
    """
    Считает score для провода: сколько операций он проходит в каждом пространстве.
    Возвращает (total, score_lp, score_lpfa, score_hv).
    total = score_lp + score_lpfa + score_hv
    """
    s_lp   = sum(wire_passes_operation(mid_gage, mid_len, op) for op in filters_dict['LEAD PREP'])
    s_lpfa = sum(wire_passes_operation(mid_gage, mid_len, op) for op in filters_dict['LEAD PREP FA'])
    s_hv   = sum(wire_passes_operation(mid_gage, mid_len, op) for op in filters_dict['High Voltage'])
    return s_lp + s_lpfa + s_hv, s_lp, s_lpfa, s_hv


# ─── ФИЧИ ДЛЯ НЕЙРОСЕТИ (CUTTING) ────────────────────────────────────────────

def create_cutting_features(df, group_mapping=None, group_global_mean=None):
    df = df.copy()

    df['mid_gage'] = (df['min_gage'] + df['max_gage']) / 2
    df['mid_len']  = (df['min_len']  + df['max_len'])  / 2
    df['gage_range'] = df['max_gage'] - df['min_gage']
    df['len_range']  = df['max_len']  - df['min_len']

    df['mid_gage_log']  = np.log1p(df['mid_gage'].clip(lower=0))
    df['mid_len_log']   = np.log1p(df['mid_len'].clip(lower=0))
    df['min_gage_log']  = np.log1p(df['min_gage'].clip(lower=0))
    df['max_gage_log']  = np.log1p(df['max_gage'].clip(lower=0))
    df['min_len_log']   = np.log1p(df['min_len'].clip(lower=0))
    df['max_len_log']   = np.log1p(df['max_len'].clip(lower=0))

    df['mid_gage_sqrt'] = np.sqrt(df['mid_gage'].clip(lower=0))
    df['mid_len_sqrt']  = np.sqrt(df['mid_len'].clip(lower=0))
    df['max_gage_sqrt'] = np.sqrt(df['max_gage'].clip(lower=0))
    df['max_len_sqrt']  = np.sqrt(df['max_len'].clip(lower=0))

    df['gage_x_len']    = df['mid_gage'] * df['mid_len']
    df['gcsp_log']      = np.log1p(df['gcsp'].clip(lower=0))

    if group_mapping is not None:
        gm = group_global_mean if group_global_mean is not None else 0.0
        df['group_encoded'] = df['machine_group'].map(group_mapping).fillna(gm)
    else:
        _, mapping, gm = _encode_machine_group(df)
        df['group_encoded'] = df['machine_group'].map(mapping).fillna(gm)

    return df


FEATURE_COLS_CUTTING = [
    'group_encoded',
    'min_gage', 'max_gage', 'min_len', 'max_len',
    'mid_gage', 'mid_len',
    'gage_range', 'len_range',
    'mid_gage_log', 'mid_len_log',
    'min_gage_log', 'max_gage_log',
    'min_len_log',  'max_len_log',
    'mid_gage_sqrt', 'mid_len_sqrt',
    'max_gage_sqrt', 'max_len_sqrt',
    'gage_x_len', 'gcsp_log',
]


# ─── ЛИСТЫ ────────────────────────────────────────────────────────────────────

# Листы для которых нейросеть предсказывает QTY/TOTAL
GENERIC_SHEETS = [
    'LEAD PREP', 'LEAD PREP FA', 'High Voltage',
]

# Листы где просто ставим QTY=1 (без модели)
SIMPLE_QTY_SHEETS = [
    'FA Conns and wires', 'FA Taping', 'FA Miscellaneos',
]


# ─── ОСТАЛЬНЫЕ ЛИСТЫ (TOTAL GCSD и generic) ──────────────────────────────────

_GCSD_SKIP_LABELS = {'total', 'per harness', 'part number', 'total time',
                     'wires number', '(min)', '%', 'summary', 'high voltage'}


def _gcsd_skip(label):
    s = label.lower().strip()
    if not s or len(s) < 2:
        return True
    for kw in _GCSD_SKIP_LABELS:
        if kw in s:
            return True
    return False


def load_and_prepare_total_gcsd(file_path, sheet_name='TOTAL values GCSD'):
    xl = pd.ExcelFile(file_path)
    df_original = xl.parse(sheet_name, header=None)
    data = []
    BLOCKS = [(1, 2, 3, 4), (9, 10, 11, 12)]

    for i in range(len(df_original)):
        row = df_original.iloc[i].values
        for label_col, gcsd_col, adj_col, plant_col in BLOCKS:
            def _c(col, row=row):
                return row[col] if col < len(row) else None
            label_raw = _c(label_col)
            try:
                if pd.isna(label_raw): continue
            except (TypeError, ValueError): pass
            label = str(label_raw).strip()
            if _gcsd_skip(label): continue
            gcsd  = _to_float(_c(gcsd_col))
            adj   = _to_float(_c(adj_col))
            plant = _to_float(_c(plant_col))
            if gcsd is None or gcsd <= 0: continue
            if adj  is None or adj  <= 0: continue
            data.append({'row_index': i, 'PART NUMBER': label,
                         'GCSD': gcsd, 'ADJ.': adj,
                         'PLANT STANDARD': plant, 'plant_col': plant_col})

    df_clean = pd.DataFrame(data) if data else pd.DataFrame(
        columns=['row_index', 'PART NUMBER', 'GCSD', 'ADJ.', 'PLANT STANDARD', 'plant_col'])
    return df_clean, df_original


def create_features_total_gcsd(df):
    df = df.copy()
    df['GCSD_log']        = np.log1p(df['GCSD'])
    df['ADJ_log']         = np.log1p(df['ADJ.'])
    df['GCSD_x_ADJ']      = df['GCSD'] * df['ADJ.']
    df['GCSD_sqrt']       = np.sqrt(df['GCSD'])
    df['ADJ_sqrt']        = np.sqrt(df['ADJ.'])
    df['PART_LEN']        = df['PART NUMBER'].astype(str).str.len()
    df['PART_DIGITS']     = df['PART NUMBER'].astype(str).str.count(r'\d')
    df['PART_HAS_LETTER'] = df['PART NUMBER'].astype(str).str.contains(r'[A-Za-z]').astype(int)
    df['PLANT STANDARD']  = df['PLANT STANDARD'].fillna(df['GCSD'] * df['ADJ.'])
    return df


FEATURE_COLS_TOTAL_GCSD = [
    'GCSD', 'ADJ.',
    'GCSD_log', 'ADJ_log',
    'GCSD_x_ADJ', 'GCSD_sqrt', 'ADJ_sqrt',
    'PART_LEN', 'PART_DIGITS', 'PART_HAS_LETTER',
]


def _detect_cols_from_header(row_vals, prev_gcsp, prev_qty, prev_total):
    new_gcsp = _col_index(row_vals, 'global sec')
    gcsp_col = new_gcsp if new_gcsp is not None else prev_gcsp
    new_qty  = _col_index(row_vals, 'qty', 'quantity')
    qty_col  = new_qty if new_qty is not None else prev_qty
    candidates = [j for j, v in enumerate(row_vals)
                  if v is not None and 'total' in str(v).lower()
                  and 'sub' not in str(v).lower()]
    total_col = candidates[-1] if candidates else prev_total
    return gcsp_col, qty_col, total_col


def load_sheet_dynamic(file_path, sheet_name):
    xl = pd.ExcelFile(file_path)
    df_raw = xl.parse(sheet_name, header=None)
    records  = []
    gcsp_col = qty_col = total_col = None

    for i, row in df_raw.iterrows():
        vals = list(row.values)
        if _is_ne(vals): continue
        if _is_header(vals):
            gcsp_col, qty_col, total_col = _detect_cols_from_header(
                vals, gcsp_col, qty_col, total_col)
            continue
        if gcsp_col is None or qty_col is None or total_col is None: continue

        gcsp  = _to_float(vals[gcsp_col]  if gcsp_col  < len(vals) else None)
        total = _to_float(vals[total_col] if total_col < len(vals) else None)
        if gcsp is None or gcsp <= 0: continue
        if total is None: continue

        qty_raw = vals[qty_col] if qty_col < len(vals) else None
        qty = _to_float(qty_raw)
        if qty is None: qty = 0.0

        cat = ''
        for v in vals:
            if v is not None:
                f = _to_float(v)
                if f is None:
                    s = str(v).strip()
                    if s and s.lower() != 'nan':
                        cat = s
                        break

        records.append({'row_index': i, 'GCSP': gcsp, 'QTY': qty, 'TOTAL': total,
                        'CATEGORY': cat, 'qty_col': qty_col, 'total_col': total_col})

    return pd.DataFrame(records), df_raw


def create_features(df):
    df = df.copy()
    df['GCSP_log']        = np.log1p(df['GCSP'])
    df['QTY_log']         = np.log1p(df['QTY'].clip(lower=0))
    df['GCSP_x_QTY']      = df['GCSP'] * df['QTY']
    df['GCSP_sqrt']       = np.sqrt(df['GCSP'].clip(lower=0))
    df['QTY_sqrt']        = np.sqrt(df['QTY'].clip(lower=0))
    df['CAT_LEN']         = df['CATEGORY'].str.len().fillna(0)
    df['CAT_DIGITS']      = df['CATEGORY'].str.count(r'\d').fillna(0)
    df['CAT_HAS_LETTERS'] = df['CATEGORY'].str.contains(r'[A-Za-z]', na=False).astype(int)
    return df


FEATURE_COLS = [
    'GCSP', 'QTY', 'GCSP_log', 'QTY_log',
    'GCSP_x_QTY', 'GCSP_sqrt', 'QTY_sqrt',
    'CAT_LEN', 'CAT_DIGITS', 'CAT_HAS_LETTERS',
]